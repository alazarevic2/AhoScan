# Gets the plain text out of an uploaded resume file

from io import BytesIO

from docx import Document
from pypdf import PdfReader


# Take a file's name and its raw bytes, and return its text
def extract_text(filename, data):
    ext = filename.rsplit(".", 1)[-1].lower()

    if ext == "txt":
        return data.decode("utf-8", errors="replace")

    if ext == "pdf":
        reader = PdfReader(BytesIO(data))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)

    if ext == "docx":
        doc = Document(BytesIO(data))
        lines = [p.text for p in doc.paragraphs]
        # Many resumes are laid out in tables, so read the text in those too
        for table in doc.tables:
            for row in table.rows:
                lines.extend(cell.text for cell in row.cells)
        return "\n".join(lines)

    raise ValueError(f"{filename}: unsupported file type (use PDF, DOCX or TXT)")
